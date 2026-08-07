from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# ── Colors ──
PRIMARY = RGBColor(0x1A, 0x56, 0xDB)      # 深蓝
SECONDARY = RGBColor(0x25, 0x91, 0xD4)     # 中蓝
ACCENT = RGBColor(0x10, 0xB9, 0x81)        # 绿色
WARM = RGBColor(0xF5, 0x9E, 0x0B)          # 橙黄
DARK = RGBColor(0x1E, 0x29, 0x3B)          # 深灰
LIGHT = RGBColor(0xF8, 0xFA, 0xFC)         # 浅灰白
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x94, 0xA3, 0xB8)
RED = RGBColor(0xEF, 0x44, 0x44)
PHASE_COLORS = [
    RGBColor(0x3B, 0x82, 0xF6),
    RGBColor(0x8B, 0x5C, 0xF6),
    RGBColor(0x10, 0xB9, 0x81),
    RGBColor(0xF5, 0x9E, 0x0B),
    RGBColor(0xEF, 0x44, 0x44),
    RGBColor(0x06, 0xB6, 0xD4),
]

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_bg(slide, color=LIGHT):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title_bar(slide, title_text, subtitle_text=""):
    """Dark header bar at top"""
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.3))
    bar.fill.solid()
    bar.fill.fore_color.rgb = DARK
    bar.line.fill.background()

    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.15), Inches(11), Inches(0.7))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = WHITE

    if subtitle_text:
        txBox2 = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11), Inches(0.4))
        tf2 = txBox2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = subtitle_text
        p2.font.size = Pt(14)
        p2.font.color.rgb = GRAY


def add_rounded_box(slide, left, top, width, height, color, opacity=1.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def add_text_box(slide, left, top, width, height, text, font_size=12, bold=False, color=DARK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return txBox


def add_multiline_box(slide, left, top, width, height, lines, default_size=11, default_color=DARK):
    """lines: list of (text, font_size, bold, color) or just str"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        if isinstance(line, str):
            text, fs, b, c = line, default_size, False, default_color
        else:
            text = line[0]
            fs = line[1] if len(line) > 1 else default_size
            b = line[2] if len(line) > 2 else False
            c = line[3] if len(line) > 3 else default_color

        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = text
        p.font.size = Pt(fs)
        p.font.bold = b
        p.font.color.rgb = c
        p.space_after = Pt(4)
    return txBox


def add_arrow(slide, start_left, start_top, end_left, end_top):
    connector = slide.shapes.add_connector(
        1, start_left, start_top, end_left, end_top
    )
    connector.line.color.rgb = GRAY
    connector.line.width = Pt(1.5)
    return connector


def add_circle_badge(slide, left, top, size, text, color, font_size=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = text
    if font_size:
        p.font.size = Pt(font_size)
    else:
        p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    return shape


# ================================================================
# SLIDE 1 — Cover
# ================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
add_bg(slide, DARK)

# Decorative bar
bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.3), prs.slide_height)
bar.fill.solid()
bar.fill.fore_color.rgb = PRIMARY
bar.line.fill.background()

# Title
add_text_box(slide, Inches(2), Inches(1.8), Inches(10), Inches(1), "Cell Vision", font_size=52, bold=True, color=WHITE)
add_text_box(slide, Inches(2), Inches(2.8), Inches(10), Inches(0.8), "单细胞识别模型  ·  识别机制说明", font_size=24, color=GRAY)

# Subtitle
add_text_box(slide, Inches(2), Inches(4.2), Inches(10), Inches(1.2),
    "基于深度学习的96孔板显微图像自动分析系统\n区分活细胞与杂质碎片，支持多时间点追踪",
    font_size=16, color=RGBColor(0xCB, 0xD5, 0xE1))

# Bottom bar
bar2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(6.9), prs.slide_width, Inches(0.6))
bar2.fill.solid()
bar2.fill.fore_color.rgb = PRIMARY
bar2.line.fill.background()

# ================================================================
# SLIDE 2 — Pipeline Overview
# ================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title_bar(slide, "整体流程概览", "从原始图像到细胞识别结果的完整 Pipeline")

phases = [
    ("1", "数据准备", "审计硬件\n构建索引\n划分数据集", 0),
    ("2", "候选生成", "多尺度峰值检测\n孔壁感知分区\n伪标签生成", 1),
    ("3", "V1弱监督", "TinyUNet训练\n教学标注\n多细胞分类", 2),
    ("4", "V2核心模型", "SeededInstanceUNet\nTemporalEvidenceNet\n实例分割", 3),
    ("5", "时间推断", "掩膜合并\n时间证据调整\n孔级筛选", 4),
    ("6", "审核与迭代", "Web审核UI\n人工修正\n模型再训练", 5),
]

start_x = Inches(0.4)
box_w = Inches(2.0)
box_h = Inches(2.8)
gap = Inches(0.15)
y_phase = Inches(1.8)
y_icon = Inches(2.8)
y_desc = Inches(3.4)
arrow_y = y_phase + box_h / 2

for i, (num, title, desc, ci) in enumerate(phases):
    x = start_x + i * (box_w + gap)

    # Phase number circle
    add_circle_badge(slide, x + Inches(0.75), y_phase, Inches(0.5), num, PHASE_COLORS[ci], font_size=14)

    # Title
    add_text_box(slide, x, y_icon, box_w, Inches(0.5), title, font_size=14, bold=True, color=DARK, align=PP_ALIGN.CENTER)

    # Description box
    box = add_rounded_box(slide, x, y_desc, box_w, Inches(2.4), WHITE)
    box.shadow.inherit = False

    # Left accent bar on box
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y_desc, Inches(0.08), Inches(2.4))
    accent.fill.solid()
    accent.fill.fore_color.rgb = PHASE_COLORS[ci]
    accent.line.fill.background()

    add_multiline_box(slide, x + Inches(0.2), y_desc + Inches(0.15), box_w - Inches(0.3), Inches(2.0),
        desc.split("\n"), default_size=10, default_color=DARK)

    # Arrow
    if i < len(phases) - 1:
        arrow_shape = slide.shapes.add_shape(
            MSO_SHAPE.RIGHT_ARROW,
            x + box_w,
            arrow_y,
            gap,
            Inches(0.2),
        )
        arrow_shape.fill.solid()
        arrow_shape.fill.fore_color.rgb = GRAY
        arrow_shape.line.fill.background()

# Bottom note
add_text_box(slide, Inches(0.8), Inches(6.5), Inches(12), Inches(0.4),
    "核心设计理念：「弱监督 → 强监督」渐进式训练  ·  「空间分割 + 时间证据」联合推理",
    font_size=12, color=GRAY, align=PP_ALIGN.CENTER)

# ================================================================
# SLIDE 3 — Candidate Detection
# ================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title_bar(slide, "候选检测：高召回率峰值检测", "dense_candidates.py — 确保不遗漏任何潜在细胞")

# Left column — Algorithm
left_x = Inches(0.6)
add_text_box(slide, left_x, Inches(1.6), Inches(5.8), Inches(0.4), "算法步骤", font_size=18, bold=True, color=PRIMARY)

steps = [
    ("① 多尺度响应图", "fine(σ=1.2), background(σ=12), medium(σ=4)\nresponse = max(背景−细尺度, 0.7×|中−细|)"),
    ("② 动态孔壁检测", "径向梯度剖面分析 → 找到壁内缘位置\n每张图独立计算, 适应拍摄偏移"),
    ("③ 三分区峰值筛选", "孔内(78%) + 壁缓冲(16%) + 壁救援(6%)\n百分位阈值 + 8×8网格覆盖 + 绝对响应底线"),
    ("④ 极坐标壁残差", "沿孔壁环形展开 → 切向平滑减除背景\n→ 恢复壁重叠的细胞 (突破性创新)"),
    ("⑤ 手动锚点注入", "人工审核标注的细胞位置强制保留\nCF分量去重 (cKDTree 距离门控)"),
]

for i, (title, desc) in enumerate(steps):
    y = Inches(2.15) + i * Inches(0.95)
    # Step number badge
    badge = add_circle_badge(slide, left_x, y + Inches(0.05), Inches(0.35), str(i+1), PRIMARY, font_size=10)
    add_text_box(slide, left_x + Inches(0.5), y, Inches(5.5), Inches(0.3), title, font_size=12, bold=True, color=DARK)
    add_text_box(slide, left_x + Inches(0.5), y + Inches(0.28), Inches(5.5), Inches(0.55), desc, font_size=10, color=GRAY)

# Right column — Key diagram
right_x = Inches(6.8)
add_text_box(slide, right_x, Inches(1.6), Inches(6), Inches(0.4), "三分区策略示意", font_size=18, bold=True, color=PRIMARY)

# Draw well diagram
cx = right_x + Inches(2.8)
cy = Inches(4.5)
outer_r = Inches(2.4)
inner_r = Inches(1.7)
buf_r = Inches(1.95)
rescue_inner = Inches(2.05)

# Outer zone (wall rescue)
o = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - outer_r, cy - outer_r, outer_r*2, outer_r*2)
o.fill.solid(); o.fill.fore_color.rgb = RGBColor(0xFE, 0xE2, 0xE2); o.line.color.rgb = GRAY; o.line.width = Pt(0.5)

# Buffer zone
b = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - rescue_inner, cy - rescue_inner, rescue_inner*2, rescue_inner*2)
b.fill.solid(); b.fill.fore_color.rgb = RGBColor(0xFE, 0xF3, 0xC7); b.line.color.rgb = GRAY; b.line.width = Pt(0.5)

# Interior zone
inn = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - buf_r, cy - buf_r, buf_r*2, buf_r*2)
inn.fill.solid(); inn.fill.fore_color.rgb = RGBColor(0xDB, 0xEA, 0xFE); inn.line.color.rgb = GRAY; inn.line.width = Pt(0.5)

# Inner zone
inner = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - inner_r, cy - inner_r, inner_r*2, inner_r*2)
inner.fill.solid(); inner.fill.fore_color.rgb = RGBColor(0xBF, 0xDB, 0xFE); inner.line.color.rgb = GRAY; inner.line.width = Pt(0.5)

# Legend
add_rounded_box(slide, right_x, Inches(6.0), Inches(6.2), Inches(1.0), WHITE)
legend_items = [
    (Inches(0.15), "■  孔内区域 (78%) — 主细胞区", RGBColor(0x3B, 0x82, 0xF6)),
    (Inches(2.2), "■  壁缓冲区 (16%) — 贴壁细胞", RGBColor(0xF5, 0x9E, 0x0B)),
    (Inches(4.2), "■  壁救援区 (6%) — 壁重叠细胞", RGBColor(0xEF, 0x44, 0x44)),
]
for _x, txt, clr in legend_items:
    add_text_box(slide, right_x + _x, Inches(6.2), Inches(2), Inches(0.3), txt, font_size=9, bold=True, color=clr)

# ================================================================
# SLIDE 4 — SeededInstanceUNet
# ================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title_bar(slide, "V2 核心模型：SeededInstanceUNet", "种子条件化实例分割 — 解决相邻细胞分离难题")

# Top — Architecture
add_text_box(slide, Inches(0.6), Inches(1.6), Inches(6), Inches(0.4), "网络架构", font_size=18, bold=True, color=PRIMARY)

arch_parts = [
    ("输入层  (3通道 × 96×96)", [
        "▸ 通道1: 归一化原始图像 (2%-98%归一化)",
        "▸ 通道2: 高斯种子热图 (种子处=1.0, 指数衰减)",
        "▸ 通道3: 几何壁先验 (距壁距离的梯度)",
    ]),
    ("编码器  (4级下采样)", [
        "▸ Enc1: 16 ch → Enc2: 32 ch → Enc3: 64 ch",
        "▸ Enc4: 128 ch → Bottleneck: 256 ch",
        "▸ 每级: Conv-BN-SiLU ×2 + MaxPool(k=2)",
    ]),
    ("解码器  (4级上采样 + 跳连)", [
        "▸ ConvTranspose2d 逐步恢复到原尺寸",
        "▸ Skip connection: 融合编码器对应层级特征",
        "▸ Dec1: 16 ch → 最终输出层",
    ]),
    ("输出层  (3通道 × 96×96)", [
        "▸ 通道1: 实例掩膜 — 当前选中物体的像素级分割",
        "▸ 通道2: 孔壁掩膜 — 标记哪些像素属于孔壁",
        "▸ 通道3: 物体存在概率 — 过滤假候选",
    ]),
]

for i, (title, lines) in enumerate(arch_parts):
    y = Inches(2.15) + i * Inches(1.2)
    box = add_rounded_box(slide, Inches(0.6), y, Inches(6.0), Inches(1.0), WHITE)
    box.shadow.inherit = False
    # Left accent
    acc = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), y, Inches(0.08), Inches(1.0))
    acc.fill.solid(); acc.fill.fore_color.rgb = ACCENT; acc.line.fill.background()
    add_text_box(slide, Inches(0.85), y + Inches(0.05), Inches(5.5), Inches(0.25), title, font_size=11, bold=True, color=PRIMARY)
    for j, line in enumerate(lines):
        add_text_box(slide, Inches(0.85), y + Inches(0.28) + j * Inches(0.18), Inches(5.5), Inches(0.18), line, font_size=9, color=DARK)

# Right side — how seed works
right_x = Inches(7.2)
add_text_box(slide, right_x, Inches(1.6), Inches(5.5), Inches(0.4), "\"种子\" 机制原理", font_size=18, bold=True, color=PRIMARY)

seed_steps = [
    ("候选坐标 → 高斯热图", "每个潜在的细胞中心 (x,y) → 在它周围生成高斯热图\n热图在中心=1.0, 向外指数衰减, 告诉模型\"关注这里\""),
    ("条件化分割", "3通道输入让模型学会:\n\"请分割离这个高斯热图中心最近的物体\""),
    ("相邻细胞可分离", "两个紧挨的细胞 → 两个独立种子 → 两次独立分割\n不再需要语义分割级别的像素硬分类"),
    ("壁先验辅助", "第3通道告知模型孔壁在哪 → 避免把壁边缘当成细胞轮廓\n同时独立预测壁掩膜, 用于后续的壁拒绝"),
]

for i, (title, desc) in enumerate(seed_steps):
    y = Inches(2.15) + i * Inches(1.2)
    box = add_rounded_box(slide, right_x, y, Inches(5.5), Inches(1.0), WHITE)
    box.shadow.inherit = False
    acc = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, right_x, y, Inches(0.08), Inches(1.0))
    acc.fill.solid(); acc.fill.fore_color.rgb = WARM; acc.line.fill.background()
    add_text_box(slide, right_x + Inches(0.2), y + Inches(0.05), Inches(5.1), Inches(0.25), title, font_size=11, bold=True, color=DARK)
    add_text_box(slide, right_x + Inches(0.2), y + Inches(0.3), Inches(5.1), Inches(0.7), desc, font_size=9, color=GRAY)

# Bottom highlight
highlight = add_rounded_box(slide, Inches(0.6), Inches(7.0), Inches(12.1), Inches(0.4), PRIMARY)
add_text_box(slide, Inches(0.8), Inches(7.02), Inches(11.5), Inches(0.35),
    "关键创新：传统U-Net做语义分割 → 相邻细胞掩膜粘连  |  SeededInstanceUNet → 种子条件化, 每次只分割一个细胞 → 天然实例级输出",
    font_size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

# ================================================================
# SLIDE 5 — TemporalEvidenceNet
# ================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title_bar(slide, "V2 时间证据网络：TemporalEvidenceNet", "利用活细胞会移动/变化、碎片不动的生物学先验区分真假")

# Left — Principle
add_text_box(slide, Inches(0.6), Inches(1.6), Inches(6), Inches(0.4), "核心原理", font_size=18, bold=True, color=PRIMARY)

principle_items = [
    ("生物学前提", "活细胞在培养过程中会移动、分裂、形态变化\n碎片/杂质/死细胞位置和形态几乎不变\n→ 对比同一位置 T0/T1/T2 的变化即可区分"),
    ("三帧联合编码", "每个时间点的 [原始图像 + 分割掩膜] 独立编码\n→ 3个96维嵌入向量 → 与数值特征拼接"),
    ("数值特征 (15~24维)", "▸ T0→T1, T1→T2, T0→T2 的位移向量 (6维)\n▸ 面积变化对数比 (2维)\n▸ 多重性编码 + 形态学概率 + 壁重叠 (7~16维)"),
    ("双输出头", "▸ same_object logit: 三帧是否是同一个物体\n▸ static_similarity logit: 物体有多静态 (越高越像碎片)"),
]

for i, (title, desc) in enumerate(principle_items):
    y = Inches(2.1) + i * Inches(1.1)
    box = add_rounded_box(slide, Inches(0.6), y, Inches(6.0), Inches(0.9), WHITE)
    box.shadow.inherit = False
    add_text_box(slide, Inches(0.8), y + Inches(0.05), Inches(5.5), Inches(0.22), title, font_size=11, bold=True, color=PRIMARY)
    add_text_box(slide, Inches(0.8), y + Inches(0.28), Inches(5.5), Inches(0.6), desc, font_size=9, color=DARK)

# Right — Decision logic
right_x = Inches(7.2)
add_text_box(slide, right_x, Inches(1.6), Inches(6), Inches(0.4), "概率调整决策流程", font_size=18, bold=True, color=PRIMARY)

decisions = [
    ("1", "invalid prob ≥ 60%?", "不调整\n(非有效物体)"),
    ("2", "same_object < 阈值?", "不调整\n(非同一物体)"),
    ("3", "static_similarity < 阈值?", "不调整\n(物体有变化, 更像活细胞)"),
    ("4", "细胞置信度 ≥ 90%?", "不调整\n(已经很确定)"),
    ("5", "所有条件满足", "计算 shift:\n碎片概率 ↑ 细胞概率 ↓\n(最多调整 0.30)"),
]

for i, (num, cond, result) in enumerate(decisions):
    y = Inches(2.1) + i * Inches(0.95)
    # Number
    add_circle_badge(slide, right_x, y + Inches(0.08), Inches(0.3), num, ACCENT if i < 4 else WARM, font_size=9)
    # Condition
    box = add_rounded_box(slide, right_x + Inches(0.5), y, Inches(3.2), Inches(0.4), WHITE)
    box.shadow.inherit = False
    add_text_box(slide, right_x + Inches(0.6), y + Inches(0.05), Inches(3), Inches(0.3), cond, font_size=10, bold=True, color=DARK)
    # Arrow
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, right_x + Inches(3.8), y + Inches(0.1), Inches(0.4), Inches(0.2))
    arrow.fill.solid(); arrow.fill.fore_color.rgb = GRAY; arrow.line.fill.background()
    # Result
    add_text_box(slide, right_x + Inches(4.4), y + Inches(0.03), Inches(2.5), Inches(0.45), result, font_size=9, color=GRAY if i < 4 else SECONDARY)

# Static wall detection
add_text_box(slide, right_x, Inches(6.0), Inches(6), Inches(0.3), "额外规则：静态壁伪影检测", font_size=14, bold=True, color=RED)
wall_rule = ("三帧都存在 + same > 阈值 + static > 阈值\n"
             "+ 位移 < 5px + 面积比 < 2.2 + 壁重叠 > 90%\n"
             "→ 标记为 \"static_wall_artifact\" → label = invalid")
add_text_box(slide, right_x, Inches(6.3), Inches(6), Inches(0.8), wall_rule, font_size=10, color=DARK)

# ================================================================
# SLIDE 6 — Review System
# ================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title_bar(slide, "人机协同审核与迭代闭环", "Human-in-the-loop — 从模型预测到人工确认的完整回路")

# Left — Review UI pages
add_text_box(slide, Inches(0.6), Inches(1.6), Inches(6), Inches(0.4), "Web 审核系统 (localhost:8765)", font_size=18, bold=True, color=PRIMARY)

review_pages = [
    ("📝 教学标注", "人工标注细胞/碎片/无效\n为教学分类器提供训练数据"),
    ("🔍 自动审核", "审核模型自动标注结果\n批量修正错误分类"),
    ("👥 双细胞教学", "专门标注接触双细胞\n训练 multiplicity 分类器"),
    ("📊 集成审核", "V2模型输出快速审核仪表板\n一键确认/修正/跳过"),
    ("🧪 孔级筛选", "孔级生长决策审核\n活跃/不活跃/模糊"),
]

for i, (title, desc) in enumerate(review_pages):
    y = Inches(2.1) + i * Inches(1.0)
    box = add_rounded_box(slide, Inches(0.6), y, Inches(5.8), Inches(0.8), WHITE)
    box.shadow.inherit = False
    add_text_box(slide, Inches(0.8), y + Inches(0.08), Inches(2.2), Inches(0.3), title, font_size=12, bold=True, color=PRIMARY)
    add_text_box(slide, Inches(3.0), y + Inches(0.08), Inches(3.2), Inches(0.65), desc, font_size=9, color=DARK)

# Right — Iteration loop
right_x = Inches(7.0)
add_text_box(slide, right_x, Inches(1.6), Inches(6), Inches(0.4), "迭代闭环", font_size=18, bold=True, color=PRIMARY)

loop_items = ["模型预测", "审核修正", "积累标注", "触发训练", "新模型"]
loop_colors = [PRIMARY, ACCENT, RGBColor(0x8B, 0x5C, 0xF6), WARM, SECONDARY]

for i, (item, clr) in enumerate(zip(loop_items, loop_colors)):
    angle = -90 + i * 72
    import math
    rad = math.radians(angle)
    cx2 = right_x + Inches(2.8)
    cy2 = Inches(4.5)
    lr = Inches(1.8)
    bx = cx2 + Emu(int(lr * math.cos(rad))) - Inches(0.65)
    by = cy2 + Emu(int(lr * math.sin(rad))) - Inches(0.35)
    box = add_rounded_box(slide, bx, by, Inches(1.3), Inches(0.7), clr)
    tf = box.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = item
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

# Center text
add_text_box(slide, right_x + Inches(1.8), Inches(4.2), Inches(2), Inches(0.6), "迭代\n闭环", font_size=14, bold=True, color=DARK, align=PP_ALIGN.CENTER)

# Bottom
add_text_box(slide, right_x, Inches(6.0), Inches(6), Inches(1.0),
    "▸ SQLite 数据库持久化所有审核结果\n▸ 触发条件: 新增标注累积 > 96 或 手动触发\n▸ 增量微调: 加载已有权重, 在新标注上 fine-tune\n▸ 外部验证: A12-22 板冻结, 仅评估, 不参与训练",
    font_size=10, color=DARK)

# ================================================================
# SLIDE 7 — Summary
# ================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)

add_text_box(slide, Inches(0.8), Inches(1.5), Inches(11), Inches(0.8), "总结：Cell Vision 识别机制三大支柱", font_size=30, bold=True, color=WHITE)

pillars = [
    ("🔬", "高召回候选检测", "多尺度响应 + 动态壁检测\n极坐标壁残差 + 三分区策略\n确保不遗漏任何潜在细胞"),
    ("🧠", "种子条件化实例分割", "高斯种子热图 + 壁先验\n每个候选独立分割\n天然解决相邻细胞粘连"),
    ("⏱️", "时间证据区分", "T0/T1/T2 三帧联合学习\n活细胞会移动/碎片不变\n生物学先验驱动概率调整"),
]

for i, (icon, title, desc) in enumerate(pillars):
    x = Inches(1.0) + i * Inches(4.0)
    box = add_rounded_box(slide, x, Inches(2.8), Inches(3.5), Inches(3.2), RGBColor(0x26, 0x34, 0x4D))
    add_text_box(slide, x + Inches(1.0), Inches(3.0), Inches(1.5), Inches(0.6), icon, font_size=36, align=PP_ALIGN.CENTER)
    add_text_box(slide, x + Inches(0.3), Inches(3.6), Inches(2.9), Inches(0.4), title, font_size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text_box(slide, x + Inches(0.3), Inches(4.1), Inches(2.9), Inches(1.5), desc, font_size=11, color=RGBColor(0xCB, 0xD5, 0xE1), align=PP_ALIGN.CENTER)

# Bottom bar
bar3 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(6.9), prs.slide_width, Inches(0.6))
bar3.fill.solid()
bar3.fill.fore_color.rgb = PRIMARY
bar3.line.fill.background()

# ── Save ──
output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Cell_Vision_识别机制说明.pptx")
prs.save(output_path)
print(f"Saved to: {output_path}")
