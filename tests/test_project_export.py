import json
import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from cellvision.project_server import (
    _export_project_rows,
    _project_review_filter_counts,
    _write_project_result_excel,
)


def _write_plate_results(root: Path, board_id: str, *, include_objects: bool = True) -> dict:
    artifact_root = root / board_id / "artifacts"
    gated_root = artifact_root / "gated"
    predictions_root = artifact_root / "predictions"
    gated_root.mkdir(parents=True)
    predictions_root.mkdir(parents=True)

    pd.DataFrame([
        {
            "well": "A1",
            "final_category": "single_cell_origin",
            "final_category_label": "单细胞来源",
            "day14_sheet_coverage_pct": 12.5,
        },
        {
            "well": "A2",
            "final_category": "no_obvious_growth",
            "final_category_label": "无明显生长",
            "day14_sheet_coverage_pct": 0.0,
        },
    ]).to_csv(gated_root / "plate_overview.csv", index=False)
    pd.DataFrame([
        {"well": "A1", "t0_cell_units": 9, "t0_cell_units_source": "human", "t1_cell_units": 2, "t2_cell_units": 2},
        {"well": "A2", "t0_cell_units": 3, "t1_cell_units": 0, "t2_cell_units": 0},
    ]).to_csv(predictions_root / "latest_well_screening.csv", index=False)
    database = artifact_root / "annotations" / "annotations.db"
    database.parent.mkdir(parents=True)
    with sqlite3.connect(database) as connection:
        connection.execute(
            "CREATE TABLE well_screening_reviews(well TEXT PRIMARY KEY, decision TEXT)"
        )
        connection.execute(
            "INSERT INTO well_screening_reviews(well, decision) VALUES ('A1', 'approved')"
        )
    if include_objects:
        pd.DataFrame([
            {"well": "A1", "timepoint": "T0", "screen_label": "single"},
            {"well": "A1", "timepoint": "T0", "screen_label": "touching_doublet"},
            {"well": "A1", "timepoint": "T0", "screen_label": "cluster_3plus"},
            {"well": "A1", "timepoint": "T1", "screen_label": "single"},
            {"well": "A1", "timepoint": "T1", "screen_label": "single"},
            {"well": "A1", "timepoint": "T2", "screen_label": "touching_doublet"},
            {"well": "A1", "timepoint": "T2", "screen_label": "debris"},
        ]).to_csv(predictions_root / "latest_screening_objects.csv", index=False)
    return {
        "slug": board_id.lower(),
        "board_id": board_id,
        "artifact_root": str(artifact_root),
        "gated_output_dir": str(gated_root),
    }


def test_project_export_contains_all_boards_and_weighted_counts(tmp_path: Path):
    manifest_path = tmp_path / "project.json"
    manifest = {
        "project_id": "demo",
        "project_name": "Demo project",
        "plates": [
            _write_plate_results(tmp_path, "Board-1"),
            _write_plate_results(tmp_path, "Board-2", include_objects=False),
        ],
    }
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")

    rows, warnings = _export_project_rows(manifest, manifest_path, task_name="Demo task")

    assert warnings == {}
    assert len(rows) == 4
    first = next(row for row in rows if row["板子名称"] == "Board-1" and row["孔号"] == "A1")
    assert first["任务名称"] == "Demo task"
    assert first["孔结论"] == "单细胞来源"
    assert first["人工判定"] == "合格"
    assert first["T0细胞总数"] == 9
    assert first["T1细胞总数"] == 2
    assert first["T2细胞总数"] == 2
    assert first["末点细胞覆盖率"] == 12.5
    assert first["Day2杂质数"] == 1

    fallback = next(row for row in rows if row["板子名称"] == "Board-2" and row["孔号"] == "A2")
    assert fallback["人工判定"] == ""
    assert fallback["T0细胞总数"] == 3
    assert fallback["末点细胞覆盖率"] == 0.0
    assert fallback["Day2杂质数"] == 0


def test_project_export_writes_readable_workbook(tmp_path: Path):
    manifest_path = tmp_path / "project.json"
    manifest = {
        "project_id": "demo",
        "project_name": "Demo project",
        "plates": [_write_plate_results(tmp_path, "Board-1")],
    }
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    destination = tmp_path / "exports" / "demo.xlsx"

    summary = _write_project_result_excel(manifest, manifest_path, destination)

    assert summary["row_count"] == 2
    workbook = load_workbook(destination, data_only=True)
    sheet = workbook["检测结果"]
    assert sheet.max_row == 3
    assert sheet.max_column == 10
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:J3"
    assert [cell.value for cell in sheet[1]] == [
        "任务名称", "板子名称", "孔号", "孔结论", "人工判定",
        "T0细胞总数", "T1细胞总数", "T2细胞总数",
        "末点细胞覆盖率", "Day2杂质数",
    ]


def test_project_review_filter_counts_use_strict_coverage_and_debris_limits(tmp_path: Path):
    manifest_path = tmp_path / "project.json"
    manifest = {
        "project_id": "demo",
        "plates": [_write_plate_results(tmp_path, "Board-1")],
    }

    matching = _project_review_filter_counts(
        manifest,
        manifest_path,
        coverage_min=10,
        debris_max=2,
        day2_cells_min=2,
        day2_cells_max=2,
    )
    coverage_boundary = _project_review_filter_counts(
        manifest, manifest_path, coverage_min=12.5
    )
    debris_boundary = _project_review_filter_counts(
        manifest, manifest_path, debris_max=1
    )

    assert matching["matching_well_count"] == 1
    assert matching["reviewable_well_count"] == 1
    assert matching["plates"][0]["matching_well_count"] == 1
    assert coverage_boundary["matching_well_count"] == 0
    assert debris_boundary["matching_well_count"] == 0
